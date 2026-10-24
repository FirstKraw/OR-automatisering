import io
from datetime import date
from html import escape

from .store import money

def invoice_pdf(snapshot, reference, created, replaces=None):
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_RIGHT
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, KeepTogether

    stream = io.BytesIO()
    doc = SimpleDocTemplate(stream, pagesize=(210 * mm, 297 * mm), leftMargin=20 * mm, rightMargin=20 * mm, topMargin=18 * mm, bottomMargin=19 * mm, title=reference, author=snapshot["settings"]["sender"])
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="Note", parent=styles["Normal"], fontSize=9, leading=13, textColor=colors.HexColor("#53605b")))
    styles.add(ParagraphStyle(name="Sum", parent=styles["Heading2"], alignment=TA_RIGHT, textColor=colors.HexColor("#193d31")))
    styles.add(ParagraphStyle(name="Sender", parent=styles["Heading2"], fontSize=17, leading=21, textColor=colors.HexColor("#193d31"), spaceAfter=7))
    styles.add(ParagraphStyle(name="HeaderRight", parent=styles["Normal"], alignment=TA_RIGHT, fontSize=10, leading=15, textColor=colors.HexColor("#193d31")))
    p = lambda text, style="Normal": Paragraph(escape(str(text)).replace("\n", "<br/>"), styles[style])
    settings, resident = snapshot["settings"], snapshot["resident"]
    sender_block = [p("AFZENDER", "Note"), p(settings["sender"], "Sender")]
    sender_block += [p(settings[key], "Note") for key in ("address", "email") if settings.get(key)]
    if settings.get("iban"):
        sender_block.append(p("Rekening: " + settings["iban"], "Note"))
    period = f"{snapshot['month'][5:7]}/{snapshot['month'][:4]}"
    header = Table([[sender_block, [p("MAALTIJDENAFREKENING", "HeaderRight"), p("Maand " + period, "HeaderRight"), p(reference, "HeaderRight"), p("Opgemaakt " + date.fromisoformat(created[:10]).strftime("%d/%m/%Y"), "HeaderRight")]]], colWidths=[100 * mm, 70 * mm])
    header.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"), ("LEFTPADDING", (0, 0), (-1, -1), 0), ("RIGHTPADDING", (0, 0), (-1, -1), 0), ("BOTTOMPADDING", (0, 0), (-1, -1), 13), ("LINEBELOW", (0, 0), (-1, -1), 1, colors.HexColor("#ccd6ce"))]))
    story = [header, Spacer(1, 6 * mm)]
    if replaces:
        story += [Spacer(1, 2 * mm), p("Vervangt afrekening " + replaces + ". Dit is het volledige gecorrigeerde maandbedrag, geen bijkomende aanrekening.", "Note")]
    story += [Spacer(1, 7 * mm), p("Bewoner: " + resident["name"], "Heading2"), p("Aan: " + resident["guardian"]), p(resident["address"]), p(resident["email"]), Spacer(1, 7 * mm)]
    rows = [[p("Omschrijving"), p("Aantal"), p("Tarief excl. btw"), p("Bedrag excl. btw")],
            [p("Gegeten maaltijden"), p(snapshot["count"]), p(money(990)), p(money(snapshot["subtotal"]))]]
    table = Table(rows, colWidths=[69 * mm, 19 * mm, 39 * mm, 43 * mm])
    table.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#edf2ed")), ("VALIGN", (0, 0), (-1, -1), "TOP"), ("TOPPADDING", (0, 0), (-1, -1), 10), ("BOTTOMPADDING", (0, 0), (-1, -1), 10), ("LINEBELOW", (0, -1), (-1, -1), .5, colors.HexColor("#ccd6ce"))]))
    story += [table, Spacer(1, 5 * mm), p("Subtotaal: " + money(snapshot["subtotal"]), "Sum"), p("Btw 6%: " + money(snapshot["vat"]), "Sum"), p("Te betalen: " + money(snapshot["total"]), "Sum"), Spacer(1, 6 * mm),
              KeepTogether([p("Betaalinstructies", "Heading2"), p("Rekening: " + settings["iban"]), p("Referentie: " + reference), p(settings["payment_note"])]), Spacer(1, 6 * mm), p("Maaltijdendatums", "Heading2"),
              p(", ".join(date.fromisoformat(d).strftime("%d/%m/%Y") for d in snapshot["dates"]) or "Geen maaltijden na correctie.")]
    def footer(canvas, document):
        canvas.setFont("Helvetica", 8)
        canvas.setFillColor(colors.HexColor("#53605b"))
        footer_reference = reference
        while canvas.stringWidth(footer_reference, "Helvetica", 8) > 135 * mm:
            footer_reference = footer_reference[:-4] + "..."
        canvas.drawString(20 * mm, 11 * mm, footer_reference)
        canvas.drawRightString(190 * mm, 11 * mm, f"Pagina {document.page}")
    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return stream.getvalue()

def month_xlsx(state):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.page import PageMargins
    wb = Workbook()
    ws = wb.active
    ws.title = "Maaltijden"
    # Hidden registrations must remain visible in the exported administration.
    dates = state.get("export_dates", state["dates"])
    meals = {(x["resident_id"], x["day"]) for x in state["meals"]}
    staff = {x["day"]: x["quantity"] for x in state["staff"]}
    people = [r for r in state["residents"] if r["active"] or any(x[0] == r["id"] for x in meals)]
    last = len(dates) + 2
    ws.append(["Maaltijden Oranje — " + state["month"]])
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last)
    ws.append(["Bewoner", *[date.fromisoformat(d) for d in dates], "Totaal"])
    for person in people:
        ws.append([person["name"], *[1 if (person["id"], d) in meals else 0 for d in dates], sum((person["id"], d) in meals for d in dates)])
    ws.append(["Begeleiders", *[staff.get(d, 0) for d in dates], sum(staff.values())])
    ws.append(["Totaal", *[sum(x[1] == d for x in meals) + staff.get(d, 0) for d in dates], len(meals) + sum(staff.values())])
    for row in ws:
        for cell in row:
            cell.font = Font(name="Calibri", size=11, color="20372F")
            cell.alignment = Alignment(vertical="center", horizontal="left" if cell.column == 1 else "center")
    for r in (1, 2, ws.max_row):
        ws.row_dimensions[r].height = 28
        for cell in ws[r]:
            cell.fill = PatternFill("solid", fgColor="E9EFE9" if r != 1 else "214E3B")
            cell.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF" if r == 1 else "20372F")
    for c in range(2, last):
        ws.cell(2, c).number_format = "ddd dd/mm"
        ws.column_dimensions[get_column_letter(c)].width = 12
        for r in range(3, ws.max_row + 1):
            ws.cell(r, c).number_format = '0;-0;""'
    ws.column_dimensions["A"].width = 29
    ws.column_dimensions[get_column_letter(last)].width = 12
    ws.freeze_panes = "B3"
    ws.print_title_rows = "1:2"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.page_margins = PageMargins(left=.25, right=.25, top=.4, bottom=.4)
    overview = wb.create_sheet("Begeleiders Oranje")
    overview.append(["Begeleidersmaaltijden — " + state["month"]])
    overview.append(["Datum", "Aantal besteld"])
    for d in dates:
        overview.append([date.fromisoformat(d), staff.get(d, 0)])
        overview.cell(overview.max_row, 1).number_format = "dd/mm/yyyy"
    overview.append(["Totaal", sum(staff.values())])
    overview.append(["Tarief en doorrekening nog niet vastgelegd."])
    overview.column_dimensions["A"].width = 52
    overview.column_dimensions["B"].width = 18
    overview.freeze_panes = "A3"
    for r in (1, 2, len(dates) + 3):
        for cell in overview[r]:
            cell.font = Font(name="Calibri", bold=True, color="20372F")
            cell.fill = PatternFill("solid", fgColor="E9EFE9")
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
